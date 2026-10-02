from __future__ import annotations

from datetime import datetime

import pytest

from openpyxl import Workbook
from openpyxl.comments import Comment


def test_default_parser_start_end_and_fix_key():
    from backend.excel.parsers import DefaultParser

    wb = Workbook()
    ws = wb.active
    ws.title = "sheet"
    # place an empty first row
    ws.append([None, None])
    # first non-empty row should be row 2
    ws.append(["Header", "Value"])
    # an empty row after data to terminate
    ws.append([None, None])

    start = DefaultParser.delineate_start_row(ws, start_row=1)
    assert start == 2
    end = DefaultParser.delineate_end_row(ws, start_row=1)
    # The first empty row is the first row in this sheet
    assert end == 1

    # fix_key basic transformations
    assert DefaultParser.fix_key("Some Key: (ignore)") == "some_key"
    assert DefaultParser.fix_key("comments") == "comment"


def test_default_keyvalue_parser_parsed_info_and_skip_merged():
    from backend.excel.parsers import DefaultKEYVALUEParser

    wb = Workbook()
    ws = wb.active
    ws.title = "kv"
    # first row empty -> start computed
    ws.append([None, None])
    ws.append(["Name", "Alice"])
    ws.append(["Comment", "hello"])
    # merged row - should be skipped
    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=2)
    # put something in merged cell region
    ws.cell(5, 1, "Merged")

    parser = DefaultKEYVALUEParser(worksheet=ws)
    pairs = list(parser.parsed_info)
    # Expect keys normalized to snake_case and comment->comment
    keys = [k for k, v in pairs]
    assert "name" in keys
    assert "comment" in keys


def test_default_table_parser_parsed_info_and_drop_empty_column():
    from backend.excel.parsers import DefaultTABLEParser

    wb = Workbook()
    ws = wb.active
    ws.title = "table"
    # header row
    ws.append(["Col A", "EmptyCol", "Col C"])
    ws.append([1, None, "x"])
    ws.append([2, None, "y"])

    parser = DefaultTABLEParser(worksheet=ws)
    rows = list(parser.parsed_info)
    assert len(rows) == 2
    # Ensure EmptyCol dropped
    assert all("emptycol" not in r for r in rows)
    assert rows[0]["col_a"] == 1


@pytest.mark.parametrize(
    "sample_id, expected",
    [
        (None, (None, 0)),
        ("", ("BLANK", 0)),
        ("ATCC123", ("ATCC123", 1)),
        ("negcontrol-1", ("negcontrol-1", -1)),
        ("pbs-", ("PBS-S1", -1)),
        ("pbs-foo", ("pbs-foo", -1)),
    ],
)
def test_clientsubmission_sample_check_sample_id(sample_id, expected):
    from backend.excel.parsers.clientsubmission_parser import ClientSubmissionSampleParser

    out = ClientSubmissionSampleParser.check_sample_id(sample_id, submitter_id="S1")
    assert out == expected


def test_copy_sheet_attributes_transfers_basic_properties():
    from backend.excel import copy_sheet_attributes

    wb = Workbook()
    src = wb.active
    src.title = "src"
    tgt = wb.create_sheet("tgt")

    # set a couple of attributes that are copied
    src.sheet_format.defaultColWidth = 12
    src.page_margins.left = 0.5
    # set a column dimension
    src.column_dimensions["A"].width = 20
    src.column_dimensions["A"].hidden = True

    copy_sheet_attributes(src, tgt)

    assert tgt.sheet_format.defaultColWidth == 12
    assert tgt.page_margins.left == 0.5
    assert tgt.column_dimensions["A"].width == 20
    assert tgt.column_dimensions["A"].hidden is True
